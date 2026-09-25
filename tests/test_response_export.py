"""Tests for post-response Excel / Word export helpers and tools."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from oracle_cpq_mcp.core.config import CPQProfile, CredentialSet
from oracle_cpq_mcp.exporters.chat_document import build_docx_from_tables
from oracle_cpq_mcp.exporters.records_excel import build_multi_sheet_workbook
from oracle_cpq_mcp.exporters.response_export import (
    MAX_EXPORT_SHEETS,
    validate_sheets_payload,
    write_export_bytes,
)
from oracle_cpq_mcp.security.context import reset_session_tool_calls
from oracle_cpq_mcp.security.rate_limit import reset_rate_limits
from oracle_cpq_mcp.security.replay import reset_replay_store
from oracle_cpq_mcp.security.settings import SecuritySettings
from oracle_cpq_mcp.tools._register import configure_security
from oracle_cpq_mcp.tools.response_export import register_response_export_tools

try:
    import docx  # noqa: F401

    HAS_DOCX = True
except ImportError:  # pragma: no cover
    HAS_DOCX = False


class FakeMcp:
    def __init__(self) -> None:
        self.tools: dict[str, Any] = {}

    def tool(self, **_kwargs: Any):
        def decorator(fn: Any) -> Any:
            self.tools[fn.__name__] = fn
            return fn

        return decorator


class FakeClient:
    def __init__(self, profile: CPQProfile) -> None:
        self.profile = profile


def _profile(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> CPQProfile:
    monkeypatch.setenv("CPQ_LOCAL_DATA_DIR", str(tmp_path / "data"))
    cfg = tmp_path / ".config"
    cfg.mkdir()
    (cfg / "demo.env").write_text(
        "CUSTOMER_NAME=Demo\nDEFAULT_ENVIRONMENT=dev\n"
        "DEV_URL=https://dev.example.com\n"
        "DEV_USERNAME=user\nDEV_PASSWORD=secret\n"
        "POST_RESPONSE_EXPORT=ask\n",
        encoding="utf-8",
    )
    monkeypatch.setenv("CPQ_CONFIG_DIR", str(cfg))
    monkeypatch.setenv("CPQ_CUSTOMER_PROFILE", "demo")
    return CPQProfile(
        customer_name="Demo",
        customer_id="demo",
        environment="dev",
        base_url="https://dev.example.com",
        credentials=[CredentialSet(username="user", password="secret")],
        rest_version="v18",
        company_login_name="_host",
        read_only=True,
        post_response_export="ask",
    )


def _register(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> dict[str, Any]:
    reset_session_tool_calls()
    reset_rate_limits()
    reset_replay_store()
    profile = _profile(tmp_path, monkeypatch)
    configure_security(
        profile,
        SecuritySettings(
            confirmation_secret="test-secret-key-for-hmac",
            confirmation_ttl_seconds=300,
            schema_integrity_enabled=False,
            max_tool_calls_per_session=100,
            rate_limit_enabled=False,
            audit_enabled=False,
            allow_prod=False,
            max_response_bytes=2_000_000,
            replay_window_seconds=60,
            read_calls_per_minute=120,
            write_calls_per_minute=10,
            privileged_calls_per_minute=5,
        ),
    )
    mcp = FakeMcp()
    register_response_export_tools(mcp, FakeClient(profile))  # type: ignore[arg-type]
    return mcp.tools


_SHEETS = [
    {
        "name": "Alpha",
        "columns": ["id", "label"],
        "rows": [{"id": 1, "label": "one"}, {"id": 2, "label": "two"}],
    },
    {
        "name": "Beta",
        "rows": [{"code": "x", "qty": 3}],
    },
]


def test_build_multi_sheet_workbook_two_sheets() -> None:
    payload = build_multi_sheet_workbook(_SHEETS)
    workbook = load_workbook(filename=BytesIO(payload))
    assert workbook.sheetnames == ["Alpha", "Beta"]
    alpha = workbook["Alpha"]
    assert [c.value for c in alpha[1]] == ["id", "label"]
    assert alpha[2][0].value == "1"
    beta = workbook["Beta"]
    assert [c.value for c in beta[1]] == ["code", "qty"]


def test_validate_sheets_payload_caps() -> None:
    with pytest.raises(ValueError, match="non-empty"):
        validate_sheets_payload([])
    too_many = [{"name": f"S{i}", "rows": []} for i in range(MAX_EXPORT_SHEETS + 1)]
    with pytest.raises(ValueError, match="max"):
        validate_sheets_payload(too_many)


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_build_docx_from_tables() -> None:
    from docx import Document

    result = build_docx_from_tables(
        title="Sample",
        sheets=_SHEETS,
        notes="Intro line",
    )
    payload = result.payload
    assert payload[:2] == b"PK"
    document = Document(BytesIO(payload))
    texts = [p.text for p in document.paragraphs]
    assert "Sample" in texts
    assert "Intro line" in texts
    assert len(document.tables) == 2
    assert result.diagrams_embedded == 0
    assert result.diagrams_skipped == []


@pytest.mark.skipif(HAS_DOCX, reason="only when python-docx missing")
def test_build_docx_requires_dependency() -> None:
    with pytest.raises(RuntimeError, match="python-docx"):
        build_docx_from_tables(title="X", sheets=_SHEETS)


# 1x1 transparent PNG
_MINI_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01"
    b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)


def test_diagram_display_inches_caps_tall_aspect() -> None:
    from oracle_cpq_mcp.exporters import chat_document as cd

    width, height = cd._diagram_display_inches(
        100, 1000, max_width=6.5, max_height=7.5
    )
    assert height == pytest.approx(7.5)
    assert width == pytest.approx(7.5 / 10.0)
    assert width < 6.5


def test_png_pixel_size_reads_ihdr() -> None:
    from oracle_cpq_mcp.exporters import chat_document as cd

    assert cd._png_pixel_size(_MINI_PNG) == (1, 1)
    assert cd._png_pixel_size(b"not-a-png") == (0, 0)


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_build_docx_embeds_fixture_png(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from docx import Document

    monkeypatch.setenv("CPQ_CONFIG_DIR", str(tmp_path / ".config"))
    (tmp_path / ".config" / "template").mkdir(parents=True)
    png_path = tmp_path / "tmp" / "demo" / "dev" / "flow.png"
    png_path.parent.mkdir(parents=True)
    png_path.write_bytes(_MINI_PNG)

    result = build_docx_from_tables(
        title="With diagram",
        sheets=_SHEETS[:1],
        notes="Before diagrams",
        diagrams=[
            {
                "title": "Order flow",
                "image_path": str(png_path),
                "caption": "Create order path",
            }
        ],
    )
    assert result.diagrams_embedded == 1
    assert result.diagrams_skipped == []
    document = Document(BytesIO(result.payload))
    texts = [p.text for p in document.paragraphs]
    assert "Order flow" in texts
    assert "Create order path" in texts
    assert "Before diagrams" in texts
    # Title/notes/diagram heading before first table heading
    assert texts.index("Before diagrams") < texts.index("Order flow")
    assert texts.index("Order flow") < texts.index("Alpha")
    assert len(document.inline_shapes) == 1
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    picture_paras = [
        p
        for p in document.paragraphs
        if p._element.xpath(".//*[local-name()='drawing']")  # noqa: SLF001
    ]
    assert picture_paras, "expected a paragraph containing the diagram image"
    assert picture_paras[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    caption_para = next(p for p in document.paragraphs if p.text == "Create order path")
    assert caption_para.alignment == WD_ALIGN_PARAGRAPH.CENTER


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_build_docx_structured_notes() -> None:
    from docx import Document

    notes = (
        "## Executive summary\n"
        "Scope line.\n"
        "\n"
        "## Results\n"
        "- PASS 8\n"
        "* FAIL 4\n"
        "Plain trailer."
    )
    result = build_docx_from_tables(
        title="Structured notes",
        sheets=_SHEETS[:1],
        notes=notes,
    )
    document = Document(BytesIO(result.payload))
    texts = [p.text for p in document.paragraphs]
    assert "Executive summary" in texts
    assert "Results" in texts
    assert "Scope line." in texts
    assert "PASS 8" in texts or "- PASS 8" in texts
    assert "FAIL 4" in texts or "* FAIL 4" in texts
    assert "Plain trailer." in texts
    assert texts.index("Executive summary") < texts.index("Scope line.")
    assert texts.index("Results") < texts.index("Alpha")

    style_by_text = {p.text: (p.style.name if p.style else "") for p in document.paragraphs}
    assert "Heading" in style_by_text.get("Executive summary", "")
    assert "Heading" in style_by_text.get("Results", "")


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_build_docx_skips_mermaid_without_mmdc(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from docx import Document
    import oracle_cpq_mcp.exporters.chat_document as chat_document
    from oracle_cpq_mcp.exporters.mermaid_render import MermaidRenderResult

    monkeypatch.setattr(
        chat_document,
        "render_mermaid_to_png",
        lambda *_a, **_k: MermaidRenderResult(
            None, "mmdc not on PATH (install @mermaid-js/mermaid-cli or pass image_path)"
        ),
    )

    result = build_docx_from_tables(
        title="Skip mermaid",
        sheets=_SHEETS[:1],
        diagrams=[
            {
                "title": "Flowchart",
                "mermaid": "flowchart LR\n  A --> B",
            }
        ],
    )
    assert result.diagrams_embedded == 0
    assert len(result.diagrams_skipped) == 1
    assert "mmdc" in result.diagrams_skipped[0]["reason"]
    document = Document(BytesIO(result.payload))
    texts = "\n".join(p.text for p in document.paragraphs)
    assert "flowchart LR" in texts
    assert "Flowchart" in texts


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_build_docx_rejects_template_image_path(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from docx import Document

    cfg = tmp_path / ".config"
    tmpl = cfg / "template"
    tmpl.mkdir(parents=True)
    monkeypatch.setenv("CPQ_CONFIG_DIR", str(cfg))
    bad = tmpl / "sneaky.png"
    bad.write_bytes(_MINI_PNG)

    result = build_docx_from_tables(
        title="Bad path",
        sheets=_SHEETS[:1],
        diagrams=[{"title": "Nope", "image_path": str(bad)}],
    )
    assert result.diagrams_embedded == 0
    assert result.diagrams_skipped
    document = Document(BytesIO(result.payload))
    assert len(document.inline_shapes) == 0


def test_export_response_word_embeds_png(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    if not HAS_DOCX:
        pytest.skip("python-docx not installed")
    tools = _register(tmp_path, monkeypatch)
    png_path = tmp_path / "tmp" / "demo" / "dev" / "d.png"
    png_path.parent.mkdir(parents=True)
    png_path.write_bytes(_MINI_PNG)
    result = tools["export_response_word"](
        title="Diagram export",
        sheets=_SHEETS[:1],
        notes="Hello",
        diagrams=[{"title": "D1", "image_path": str(png_path)}],
    )
    assert isinstance(result, list)
    lead = result[0]
    assert lead["status"] == "ok"
    assert lead["data"]["diagrams_embedded"] == 1
    assert lead["data"]["diagrams_skipped"] == []
    assert Path(lead["data"]["absolute_path"]).exists()


def test_offer_export_response_needs_input(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    tools = _register(tmp_path, monkeypatch)
    result = tools["offer_export_response"](title="Demo table", sheets=_SHEETS)
    assert result["status"] == "ok"
    assert result["data"]["needs_user_input"] is True
    assert any("excel" in c for c in result["data"]["choices"])


def test_export_response_excel_writes_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    tools = _register(tmp_path, monkeypatch)
    result = tools["export_response_excel"](title="Demo table", sheets=_SHEETS)
    assert isinstance(result, list)
    lead = result[0]
    assert lead["status"] == "ok"
    path = Path(lead["data"]["absolute_path"])
    assert path.exists()
    assert path.suffix == ".xlsx"
    assert "exports" in lead["data"]["path"]


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_export_response_word_writes_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    tools = _register(tmp_path, monkeypatch)
    result = tools["export_response_word"](
        title="Demo table", sheets=_SHEETS, notes="Hello"
    )
    assert isinstance(result, list)
    lead = result[0]
    assert lead["status"] == "ok"
    assert lead["data"]["uri"].startswith("file:")
    assert Path(lead["data"]["absolute_path"]).exists()


def test_set_post_response_export_writes_env(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    tools = _register(tmp_path, monkeypatch)
    result = tools["set_post_response_export"](policy="never")
    assert result["data"]["policy"] == "never"
    text = (tmp_path / ".config" / "demo.env").read_text(encoding="utf-8")
    assert "POST_RESPONSE_EXPORT=never" in text


def test_write_export_bytes_roundtrip(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    profile = _profile(tmp_path, monkeypatch)
    path = write_export_bytes(profile, "sample.xlsx", b"abc")
    assert path.read_bytes() == b"abc"


def _tbl_header_set(row: Any) -> bool:
    from docx.oxml.ns import qn

    tr = row._tr  # noqa: SLF001
    tr_pr = tr.find(qn("w:trPr"))
    if tr_pr is None:
        return False
    return tr_pr.find(qn("w:tblHeader")) is not None


def _cell_width_inches(cell: Any) -> float:
    from docx.shared import Inches

    width = cell.width
    assert width is not None
    return float(width) / float(Inches(1))


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_build_docx_applies_fixed_column_widths(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from docx import Document

    monkeypatch.setenv("CPQ_CONFIG_DIR", str(tmp_path / ".config"))
    monkeypatch.setenv("CPQ_TEMPLATE_WORK_DIR", str(tmp_path / "templates"))
    (tmp_path / ".config" / "template").mkdir(parents=True)

    result = build_docx_from_tables(
        title="Widths",
        sheets=[
            {
                "name": "Narrow",
                "columns": ["id", "label", "notes"],
                "rows": [
                    {
                        "id": "identifier-01",
                        "label": "short",
                        "notes": "notes-field-value-xx",
                    }
                ],
            }
        ],
    )
    document = Document(BytesIO(result.payload))
    assert len(document.tables) == 1
    table = document.tables[0]
    assert len(table.columns) == 3
    assert _tbl_header_set(table.rows[0])
    widths = [_cell_width_inches(c) for c in table.rows[0].cells]
    assert widths[2] > widths[0]
    header_run = table.rows[0].cells[0].paragraphs[0].runs[0]
    assert header_run.bold is True
    assert header_run.font.size is not None


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_build_docx_uses_landscape_for_wide_table(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT

    monkeypatch.setenv("CPQ_CONFIG_DIR", str(tmp_path / ".config"))
    monkeypatch.setenv("CPQ_TEMPLATE_WORK_DIR", str(tmp_path / "templates"))
    (tmp_path / ".config" / "template").mkdir(parents=True)

    cols = [f"c{i}" for i in range(6)]
    row = {c: f"v{i}" * 8 for i, c in enumerate(cols)}
    result = build_docx_from_tables(
        title="Wide",
        sheets=[{"name": "WideSheet", "columns": cols, "rows": [row]}],
    )
    document = Document(BytesIO(result.payload))
    assert any(s.orientation == WD_ORIENT.LANDSCAPE for s in document.sections)
    assert len(document.tables) == 1


@pytest.mark.skipif(not HAS_DOCX, reason="python-docx not installed")
def test_build_docx_uses_blocks_when_too_many_columns(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from docx import Document

    monkeypatch.setenv("CPQ_CONFIG_DIR", str(tmp_path / ".config"))
    monkeypatch.setenv("CPQ_TEMPLATE_WORK_DIR", str(tmp_path / "templates"))
    (tmp_path / ".config" / "template").mkdir(parents=True)

    cols = [
        "Data table",
        "Result",
        "BML file",
        "BML path",
        "Exact BMQL",
        "Expected keys",
        "Actual keys",
        "Missing keys",
    ]
    row = {
        "Data table": "Community_F44H101",
        "Result": "pass",
        "BML file": "calc.bml",
        "BML path": "data/drees/dev/bml/site/commerce/calc.bml",
        "Exact BMQL": 'bmql("select AGM from Community_F44H101 where CMHBMCUS = $x")',
        "Expected keys": "CMHBMCUS",
        "Actual keys": "CMHBMCUS, CMCPHASE",
        "Missing keys": "",
    }
    result = build_docx_from_tables(
        title="Blocks",
        sheets=[{"name": "BMQL statements", "columns": cols, "rows": [row]}],
    )
    document = Document(BytesIO(result.payload))
    # Block layout: one 2-col label/value table per record (not an 8-col grid).
    assert len(document.tables) >= 1
    assert all(len(t.columns) == 2 for t in document.tables)
    texts = [p.text for p in document.paragraphs]
    assert "Community_F44H101" in texts
